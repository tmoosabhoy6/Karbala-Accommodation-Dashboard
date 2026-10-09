"""Convert the "Krb Accomodation" xlsx (ACCOMMODATION tab) into seed SQL.

Layout of the tab: one row per room (type in A, number in B, beds in C, extra
mattresses in D) and one column per night (Gregorian date in row 3). Consecutive
nights with the same text become one stay: check-in 12:00 on the first night,
check-out 08:00 the morning after the last night (Asia/Baghdad).

Usage: python3 scripts/parse_sheet.py krb.xlsx out.sql [--since 2026-06-01]
       python3 scripts/parse_sheet.py krb.xlsx public/demo-seed.json --since 2026-09-01
A .json target writes the seed for the browser-only demo mode instead (never commit it).
"""
import datetime as dt
import re
import sys

import openpyxl

BUILDINGS = [("amatullah", "Amatullah", "AMT", 1), ("qasr", "Qasr", "QSR", 2)]
BLOCKS = [(6, 116, "amatullah"), (120, 157, "qasr")]  # rows 161-184 are an unused old Qasr list
STAFF_ROWS = (189, 205)
SHARED_ROOMS = {"amatullah:135": "female", "amatullah:219": "male"}  # Bairao / Marado rooms
STAFF_SINCE = "2026-10-01"

TYPE_NAMES = {
    "PARTITION PARDA": "Partition Parda",
    "BATHROOM INSIDE": "Bathroom Inside",
    "BATHROOM OUTSIDE": "Bathroom Outside",
    "BAIRAO ROOM": "Bairao Room",
    "MARADO ROOM / BATHROOM OUTSIDE": "Marado Room",
    "ROOM + HALL/ 1 BATHROOM": "Room + Hall",
    "ROOM + HALL / SOFA-COMEBED / 2 BATHROOMS": "Suite, 2 Bathrooms",
    "FRONT WINDOW": "Front Window",
    "2 ROOMS / 2 BATHROOMS": "2 Rooms, 2 Bathrooms",
    "NO WINDOW": "No Window",
    "NO WINDOW/LOW CEILING": "No Window, Low Ceiling",
    "WINDOWS+SOFA": "Windows + Sofa",
    "2 BHK": "2 BHK",
    "2 BEDROOMS": "2 Bedrooms",
}

BLOCKED = re.compile(r"^(faiz room|maint|main|maintenance|paint|polish and maintenance|store)$", re.I)
DROP = re.compile(r"^dm$", re.I)


def norm(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return None if s in ("", ".") else s


def beds_of(v):
    s = norm(v)
    if not s:
        return 0
    if "+" in s:
        return sum(int(p) for p in s.split("+") if p.strip().isdigit())
    try:
        return int(float(s))
    except ValueError:
        return 0


def display_number(raw, building):
    s = raw.upper()
    s = re.sub(r"^(\d+)\s*([A-Z])$", r"\1 \2", s)
    if building == "amatullah":
        s = re.sub(r"^([GA]) ?(\d)$", r"\1\2", s)
    return s


def floor_of(number, building):
    s = number[2:] if building == "qasr" and number.startswith("Q ") else number
    if s.startswith("G"):
        return "G", 0
    if s.startswith("A"):
        return "Annex", 9
    m = re.match(r"(\d)\d\d", s)
    return (m.group(1), int(m.group(1))) if m else ("?", 8)


def sort_key(number):
    m = re.search(r"(\d+)", number)
    if not m:
        return 0
    suffix = re.search(r"\d\s*([A-F])$", number)
    return int(m.group(1)) * 10 + (ord(suffix.group(1)) - 64 if suffix else 0)


def room_id(building, number):
    return f"{'amt' if building == 'amatullah' else 'qsr'}-" + re.sub(r"[^a-z0-9]", "", number.lower().removeprefix("q "))


def classify(label):
    low = label.lower().strip()
    if DROP.match(low):
        return None
    if BLOCKED.match(low):
        return ("blocked", None, label.title())
    if low == "gl":
        return ("group_leader", None, "Group Leader")
    if low in ("kg", "kgs") or re.search(r"\bkgs?\b|khidmat", low):
        return ("khidmat_guzar", None, label if low not in ("kg", "kgs") else "Khidmat Guzar")
    if re.fullmatch(r"(roti )?hr", low):
        return ("hr_mawaid", None, "Roti HR" if low.startswith("roti") else "HR Mawaid")
    ids = re.findall(r"(?<!\d)\d{3,5}(?!\d)", label)
    group = re.sub(r"(?<!\d)\d{3,5}(?!\d)", " ", label)
    group = re.sub(r"\s*/\s*", " ", group)
    group = re.sub(r"\s+", " ", group).strip(" -/") or None
    return ("tour", "/".join(ids) or None, group)


def q(v):
    return "null" if v is None else "'" + str(v).replace("'", "''") + "'"


def main(path, out, since):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["ACCOMMODATION"]
    dates = {c: ws.cell(3, c).value.date() for c in range(6, ws.max_column + 1)
             if isinstance(ws.cell(3, c).value, dt.datetime)}
    rooms, stays = {}, []

    for start, end, building in BLOCKS:
        for r in range(start, end + 1):
            raw = norm(ws.cell(r, 2).value)
            if not raw or raw == "ROOMS":
                continue
            number = display_number(raw, building)
            rid = room_id(building, number)
            if rid in rooms:  # 135 and 219 are split over several bed rows
                rooms[rid]["beds"] += beds_of(ws.cell(r, 3).value)
                rooms[rid]["extra"] += beds_of(ws.cell(r, 4).value)
            else:
                floor, fsort = floor_of(number, building)
                t = norm(ws.cell(r, 1).value)
                rooms[rid] = {
                    "id": rid, "building": building, "number": number, "floor": floor, "floor_sort": fsort,
                    "type": TYPE_NAMES.get(t.upper(), t.title()) if t else "Standard",
                    "beds": beds_of(ws.cell(r, 3).value), "extra": beds_of(ws.cell(r, 4).value),
                    "sharing": SHARED_ROOMS.get(f"{building}:{raw}", "none"), "sort": sort_key(number),
                }
            run = None
            for c in sorted(dates):
                val = norm(ws.cell(r, c).value)
                if run and val == run["label"] and dates[c] == run["last"] + dt.timedelta(days=1):
                    run["last"] = dates[c]
                    continue
                if run:
                    stays.append(run)
                run = {"room": rid, "label": val, "first": dates[c], "last": dates[c]} if val else None
            if run:
                stays.append(run)

    staff = []
    for r in range(STAFF_ROWS[0], STAFF_ROWS[1] + 1):
        raw = norm(ws.cell(r, 2).value)
        if not raw:
            continue
        number = display_number(raw, "amatullah")
        rid = room_id("amatullah", number)
        floor, fsort = floor_of(number, "amatullah")
        rooms.setdefault(rid, {"id": rid, "building": "amatullah", "number": number, "floor": floor,
                               "floor_sort": fsort, "type": "Standard", "beds": 0, "extra": 0,
                               "sharing": "none", "sort": sort_key(number)})
        staff.append((rid, norm(ws.cell(r, 4).value)))

    rows = []
    for s in stays:
        co = s["last"] + dt.timedelta(days=1)
        if co.isoformat() < since:
            continue
        c = classify(s["label"])
        if not c:
            continue
        rows.append((s["room"], c[0], c[1], c[2], s["first"].isoformat(), co.isoformat()))

    if out.endswith(".json"):
        write_json(out, rooms, rows, staff, max(dates.values()))
        print(f"{len(rooms)} rooms, {len(rows)} stays, {len(staff)} staff")
        return

    with open(out, "w") as f:
        f.write("insert into public.buildings (id, name, code, sort) values\n")
        f.write(",\n".join(f"({q(b[0])},{q(b[1])},{q(b[2])},{b[3]})" for b in BUILDINGS) + ";\n")
        f.write("insert into public.rooms (id, building_id, number, floor, floor_sort, room_type, beds, extra_beds, sharing, sort) values\n")
        f.write(",\n".join(
            f"({q(r['id'])},{q(r['building'])},{q(r['number'])},{q(r['floor'])},{r['floor_sort']},{q(r['type'])},{r['beds']},{r['extra']},{q(r['sharing'])},{r['sort']})"
            for r in rooms.values()) + ";\n")
        f.write("insert into public.stays (room_id, category, tour_id, group_name, check_in, check_out, checked_out_at, auto_checked_out, source)\n"
                "select v.room, v.cat, v.tour, v.grp, (v.ci::date + time '12:00') at time zone 'Asia/Baghdad',\n"
                " (v.co::date + time '08:00') at time zone 'Asia/Baghdad',\n"
                " case when (v.co::date + time '08:00') at time zone 'Asia/Baghdad' <= now() then (v.co::date + time '08:00') at time zone 'Asia/Baghdad' end,\n"
                " (v.co::date + time '08:00') at time zone 'Asia/Baghdad' <= now(), 'sheet'\n"
                "from (values\n")
        f.write(",\n".join(f"({q(a)},{q(b)},{q(c)},{q(d)},{q(e)},{q(g)})" for a, b, c, d, e, g in rows))
        f.write("\n) as v(room, cat, tour, grp, ci, co);\n")
        f.write("insert into public.stays (room_id, category, guest_name, check_in, source) values\n")
        f.write(",\n".join(f"({q(rid)},'staff',{q(name)},('{STAFF_SINCE}'::date + time '12:00') at time zone 'Asia/Baghdad','sheet')"
                           for rid, name in staff) + ";\n")
    print(f"{len(rooms)} rooms, {len(rows)} stays, {len(staff)} staff")


def write_json(out, rooms, rows, staff, sheet_end):
    import json
    import uuid

    tz = dt.timezone(dt.timedelta(hours=3))
    now = dt.datetime.now(dt.timezone.utc)
    iso = lambda d, hhmm: dt.datetime.combine(dt.date.fromisoformat(d), dt.time(*map(int, hhmm.split(":"))), tz).astimezone(dt.timezone.utc).isoformat()
    staff_ids = {rid for rid, _ in staff}
    base = {"guest_name": None, "phone": None, "pax": None, "notes": None, "checked_out_at": None, "auto_checked_out": False,
            "cancelled_at": None, "transferred_from": None, "transferred_to": None, "arrived_at": None, "source": "sheet",
            "created_at": now.isoformat(), "updated_at": now.isoformat()}
    stays = []
    last = (sheet_end + dt.timedelta(days=1)).isoformat()
    for room, cat, tour, grp, ci, co in rows:
        out_iso = None if co >= last else iso(co, "08:00")  # still running when the sheet ends
        done = out_iso is not None and dt.datetime.fromisoformat(out_iso) <= now
        stays.append({**base, "id": str(uuid.uuid4()), "room_id": room, "category": cat, "tour_id": tour, "group_name": grp,
                      "check_in": iso(ci, "12:00"), "check_out": out_iso, "checked_out_at": out_iso if done else None,
                      "auto_checked_out": done, "arrived_at": iso(ci, "12:00")})
    for rid, name in staff:
        stays.append({**base, "id": str(uuid.uuid4()), "room_id": rid, "category": "staff", "tour_id": None, "group_name": None,
                      "guest_name": name, "check_in": iso(STAFF_SINCE, "12:00"), "check_out": None})
    data = {
        "buildings": [{"id": b[0], "name": b[1], "code": b[2], "sort": b[3]} for b in BUILDINGS],
        "rooms": [{"id": r["id"], "building_id": r["building"], "number": r["number"], "floor": r["floor"], "floor_sort": r["floor_sort"],
                   "room_type": "Staff Room" if r["id"] in staff_ids else r["type"], "beds": 0 if r["id"] in staff_ids else r["beds"],
                   "extra_beds": 0 if r["id"] in staff_ids else r["extra"], "sharing": r["sharing"], "notes": None, "sort": r["sort"],
                   "active": True} for r in rooms.values()],
        "stays": stays, "tours": [], "activity": [],
    }
    with open(out, "w") as f:
        json.dump(data, f)


if __name__ == "__main__":
    since = sys.argv[sys.argv.index("--since") + 1] if "--since" in sys.argv else "2000-01-01"
    main(sys.argv[1], sys.argv[2], since)
